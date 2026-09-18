"""Tests para scripts/cna_extraction/sheet_resolver.py — casos normal,
'pestaña contiene número interno distinto' y 'hoja huérfana'."""
from __future__ import annotations

import openpyxl
import pytest

from scripts.cna_extraction.models import CatalogRecord
from scripts.cna_extraction.sheet_resolver import resolve_sheets


def _catalog_record(numero: int, tipo: str = "tabla", nombre: str = "x") -> CatalogRecord:
    return CatalogRecord(
        orden=None,
        factor_raw="",
        factor_num=None,
        factor_nombre="",
        caracteristica="",
        tabla_no=numero if tipo == "tabla" else None,
        grafico_no=numero if tipo == "grafico" else None,
        nombre=nombre,
        fuente="",
        responsable="",
        enlace="",
        observaciones="",
        tipo=tipo,
        numero=numero,
        id_hint=f"Tabla {numero}" if tipo == "tabla" else f"Gráfico {numero}",
    )


@pytest.fixture
def workbook():
    wb = openpyxl.Workbook()
    default = wb.active
    default.title = "Índice Tablas"

    ws1 = wb.create_sheet("Tabla 1")
    ws1.append(["Tabla NO", 1])
    ws1.append(["Nombre", "Indicador normal"])

    # Caso real: la pestaña se llama "Tabla 127" pero declara internamente 130.
    ws2 = wb.create_sheet("Tabla 127")
    ws2.append(["Tabla NO", 130])
    ws2.append(["Nombre", "Indicador con número interno distinto"])

    # Hoja huérfana: no tiene registro de catálogo correspondiente.
    ws3 = wb.create_sheet("Tabla 999")
    ws3.append(["Tabla NO", 999])
    ws3.append(["Nombre", "Sin catálogo"])

    return wb


def test_normal_match_by_declared_number(workbook):
    catalog = [_catalog_record(1)]
    resolved, unreferenced = resolve_sheets(catalog, workbook)
    match = next(r for r in resolved if r.sheet_name == "Tabla 1")
    assert match.catalog_record is not None
    assert match.catalog_record.numero == 1
    assert match.match_method == "declared_number"


def test_tab_name_mismatch_resolved_by_internal_number(workbook):
    catalog = [_catalog_record(130)]
    resolved, unreferenced = resolve_sheets(catalog, workbook)
    match = next(r for r in resolved if r.sheet_name == "Tabla 127")
    assert match.catalog_record is not None
    assert match.catalog_record.numero == 130
    assert match.declared_number == 130


def test_orphan_sheet_reported_as_unreferenced(workbook):
    catalog = [_catalog_record(1), _catalog_record(130)]
    resolved, unreferenced = resolve_sheets(catalog, workbook)
    assert "Tabla 999" in unreferenced


def test_catalog_record_without_sheet_is_reported(workbook):
    catalog = [_catalog_record(1), _catalog_record(500)]
    resolved, unreferenced = resolve_sheets(catalog, workbook)
    orphan_catalog = [r for r in resolved if r.sheet_name is None]
    assert any(r.catalog_record.numero == 500 for r in orphan_catalog)
